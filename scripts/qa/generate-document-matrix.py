#!/usr/bin/env python3
"""Generate deterministic, entirely fictional CV files for local parser QA."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BLUE = RGBColor(0x2E, 0x74, 0xB5)
GRAY = RGBColor(0x55, 0x55, 0x55)


PROFILES = {
    "pl": {
        "name": "Marta Nowak-Testowa",
        "title": "Senior Frontend Developer",
        "contact": "marta.nowak@example.invalid | +48 000 000 000 | Warszawa",
        "summary_heading": "Profil zawodowy",
        "summary": "Frontend developerka z osmioma latami doswiadczenia w produktach SaaS, wydajnosci interfejsow i systemach projektowych. Laczy decyzje techniczne z pomiarem efektu dla uzytkownikow.",
        "experience_heading": "Doswiadczenie zawodowe",
        "experience": [
            ("Senior Frontend Developer - Fikcyjna Chmura Sp. z o.o. | 2022-obecnie", [
                "Prowadzila migracje aplikacji do nowej architektury komponentow i ograniczyla czas wdrozen zespolu.",
                "Wprowadzila budzety Web Vitals, testy regresji i dashboard jakosci dla trzech produktow.",
                "Uzgadniala priorytety z product managerami, projektantami i zespolami backendowymi.",
            ]),
            ("Frontend Developer - Demo Commerce Labs | 2018-2022", [
                "Rozwijala checkout i eksperymenty produktowe w React oraz TypeScript.",
                "Uproscila biblioteke formularzy i poprawila dostepnosc kluczowych sciezek zakupowych.",
            ]),
        ],
        "skills_heading": "Umiejetnosci",
        "skills": "TypeScript, React, Next.js, testy automatyczne, web performance, accessibility, CI/CD",
        "education_heading": "Wyksztalcenie",
        "education": "Informatyka, Politechnika Przykladowa, 2014-2018",
        "languages_heading": "Jezyki",
        "languages": "Polski - ojczysty; angielski - C1",
    },
    "en": {
        "name": "Alex Example",
        "title": "Platform Engineer",
        "contact": "alex.example@example.invalid | +1 202 555 0100 | Example City",
        "summary_heading": "Professional summary",
        "summary": "Platform engineer with seven years of experience improving delivery systems, cloud reliability and developer workflows. Turns operational constraints into measurable engineering improvements.",
        "experience_heading": "Professional experience",
        "experience": [
            ("Senior Platform Engineer - Sample Systems Ltd | 2021-present", [
                "Led a deployment-platform redesign used by six product teams and introduced progressive delivery controls.",
                "Built service-level dashboards and incident review practices that reduced recurring operational failures.",
                "Partnered with security and product engineering on safe defaults for cloud infrastructure.",
            ]),
            ("DevOps Engineer - Fictional Data Works | 2017-2021", [
                "Automated infrastructure changes and standardized CI pipelines across customer-facing services.",
                "Improved recovery documentation and tested disaster-recovery procedures every quarter.",
            ]),
        ],
        "skills_heading": "Skills",
        "skills": "Kubernetes, Terraform, AWS, observability, incident response, GitHub Actions, Python",
        "education_heading": "Education",
        "education": "BSc Computer Science, Example Institute of Technology, 2013-2017",
        "languages_heading": "Languages",
        "languages": "English - fluent; Polish - intermediate",
    },
}


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, RGBColor(0x1F, 0x4D, 0x78)),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_cv_header(doc: Document, profile: dict) -> None:
    name = doc.add_paragraph()
    name.paragraph_format.space_after = Pt(4)
    run = name.add_run(profile["name"])
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(23)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run(profile["title"])
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = GRAY
    contact = doc.add_paragraph(profile["contact"])
    contact.paragraph_format.space_after = Pt(12)


def add_cv_body(doc: Document, profile: dict, use_bullets: bool = True) -> None:
    doc.add_heading(profile["summary_heading"], level=1)
    doc.add_paragraph(profile["summary"])
    doc.add_heading(profile["experience_heading"], level=1)
    for title, bullets in profile["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.add_run(title).bold = True
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet" if use_bullets else None)
    doc.add_heading(profile["skills_heading"], level=1)
    doc.add_paragraph(profile["skills"])
    doc.add_heading(profile["education_heading"], level=1)
    doc.add_paragraph(profile["education"])
    doc.add_heading(profile["languages_heading"], level=1)
    doc.add_paragraph(profile["languages"])


def save_standard_docx(path: Path, language: str) -> None:
    doc = Document()
    configure_doc(doc)
    add_cv_header(doc, PROFILES[language])
    add_cv_body(doc, PROFILES[language])
    doc.save(path)


def save_table_docx(path: Path) -> None:
    profile = PROFILES["pl"]
    doc = Document()
    configure_doc(doc)
    add_cv_header(doc, profile)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    widths = (Inches(1.875), Inches(4.625))
    rows = [
        (profile["summary_heading"], profile["summary"]),
        (profile["experience_heading"], " ".join(title + " " + " ".join(bullets) for title, bullets in profile["experience"])),
        (profile["skills_heading"], profile["skills"]),
        (profile["education_heading"], profile["education"]),
        (profile["languages_heading"], profile["languages"]),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = value
    doc.save(path)


def save_two_column_docx(path: Path) -> None:
    profile = PROFILES["en"]
    doc = Document()
    configure_doc(doc)
    add_cv_header(doc, profile)
    section = doc.sections[0]
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols_node = cols[0] if cols else OxmlElement("w:cols")
    if not cols:
        sect_pr.append(cols_node)
    cols_node.set(qn("w:num"), "2")
    cols_node.set(qn("w:space"), "720")
    add_cv_body(doc, profile)
    doc.save(path)


def save_header_contact_docx(path: Path) -> None:
    profile = PROFILES["pl"]
    doc = Document()
    configure_doc(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.text = f"{profile['name']} | {profile['contact']}"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(profile["title"])
    run.bold = True
    run.font.size = Pt(23)
    add_cv_body(doc, profile)
    doc.save(path)


def save_mixed_language_docx(path: Path) -> None:
    profile = dict(PROFILES["en"])
    profile["summary_heading"] = "Profil zawodowy / Professional summary"
    profile["summary"] = "Platform engineer z siedmioletnim doswiadczeniem. Builds reliable delivery systems and measurable developer workflows for international product teams."
    profile["skills_heading"] = "Umiejetnosci / Skills"
    profile["education_heading"] = "Wyksztalcenie / Education"
    doc = Document()
    configure_doc(doc)
    add_cv_header(doc, profile)
    add_cv_body(doc, profile)
    doc.save(path)


SHORT_TEXT = (
    "Nina Demo - Frontend Developer\n"
    "Doswiadczenie zawodowe\n"
    "Frontend Developer, Test Labs, 2024-obecnie. Tworzenie interfejsow React i testow.\n"
    "Umiejetnosci\nReact, TypeScript, CSS, Git.\n"
    "Wyksztalcenie\nInformatyka, 2020-2024."
)


def save_short_docx(path: Path) -> None:
    doc = Document()
    configure_doc(doc)
    for block in SHORT_TEXT.split("\n"):
        doc.add_paragraph(block)
    doc.save(path)


def make_scan_image(path: Path) -> None:
    image = Image.new("RGB", (1275, 1650), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(__file__).resolve().parents[2] / "src" / "assets" / "og-font.ttf"
    body = ImageFont.truetype(str(font_path), 28)
    title = ImageFont.truetype(str(font_path), 48)
    draw.text((90, 80), "Ola Fikcyjna", fill="black", font=title)
    scan_text = (
        "QA Automation Engineer\nexample@example.invalid | +48 000 000 000\n\n"
        "DOSWIADCZENIE ZAWODOWE\nQA Engineer - Demo Quality Lab | 2020-obecnie\n"
        "- Rozwijanie testow API i interfejsu\n- Utrzymanie pipeline CI\n\n"
        "UMIEJETNOSCI\nPlaywright, TypeScript, API testing, CI/CD\n\n"
        "WYKSZTALCENIE\nInformatyka, 2016-2020"
    )
    draw.multiline_text((90, 175), scan_text, fill="black", font=body, spacing=16)
    image.save(path, quality=95)


def save_image_only_docx(path: Path, scan_path: Path) -> None:
    doc = Document()
    configure_doc(doc)
    doc.add_picture(str(scan_path), width=Inches(6.5))
    doc.save(path)


def save_malformed_docx(path: Path) -> None:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", "<w:document><broken>")


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("CvTitle", parent=styles["Title"], fontName="CVFixture", fontSize=23, leading=27, textColor=colors.black, alignment=TA_LEFT, spaceAfter=4),
        "subtitle": ParagraphStyle("CvSubtitle", parent=styles["Normal"], fontName="CVFixture", fontSize=14, leading=17, textColor=colors.HexColor("#555555"), spaceAfter=8),
        "heading": ParagraphStyle("CvHeading", parent=styles["Heading1"], fontName="CVFixture", fontSize=16, leading=19, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=6),
        "body": ParagraphStyle("CvBody", parent=styles["BodyText"], fontName="CVFixture", fontSize=10.5, leading=13.2, spaceAfter=6),
        "bullet": ParagraphStyle("CvBullet", parent=styles["BodyText"], fontName="CVFixture", fontSize=10.5, leading=13.2, leftIndent=18, firstLineIndent=-9, spaceAfter=4),
    }


def pdf_story(profile: dict):
    styles = pdf_styles()
    story = [
        Paragraph(profile["name"], styles["title"]),
        Paragraph(profile["title"], styles["subtitle"]),
        Paragraph(profile["contact"], styles["body"]),
        Paragraph(profile["summary_heading"], styles["heading"]),
        Paragraph(profile["summary"], styles["body"]),
        Paragraph(profile["experience_heading"], styles["heading"]),
    ]
    for title, bullets in profile["experience"]:
        story.append(Paragraph(f"<b>{title}</b>", styles["body"]))
        story.extend(Paragraph(f"- {bullet}", styles["bullet"]) for bullet in bullets)
    story.extend([
        Paragraph(profile["skills_heading"], styles["heading"]),
        Paragraph(profile["skills"], styles["body"]),
        Paragraph(profile["education_heading"], styles["heading"]),
        Paragraph(profile["education"], styles["body"]),
        Paragraph(profile["languages_heading"], styles["heading"]),
        Paragraph(profile["languages"], styles["body"]),
    ])
    return story


def save_standard_pdf(path: Path, language: str) -> None:
    document = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    document.build(pdf_story(PROFILES[language]))


def save_multicolumn_pdf(path: Path) -> None:
    profile = PROFILES["en"]
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    pdf.setFont("CVFixture", 23)
    pdf.drawString(inch, height - inch, profile["name"])
    pdf.setFont("CVFixture", 14)
    pdf.setFillColor(colors.HexColor("#555555"))
    pdf.drawString(inch, height - inch - 26, profile["title"])
    pdf.setFillColor(colors.black)
    left_lines = [profile["summary_heading"], profile["summary"], profile["skills_heading"], profile["skills"]]
    right_lines = [profile["experience_heading"]]
    for title, bullets in profile["experience"]:
        right_lines.extend([title, *[f"- {item}" for item in bullets]])

    def draw_column(lines, x, y, max_chars):
        pdf.setFont("CVFixture", 9.5)
        for line in lines:
            words = line.split()
            rows = []
            current = ""
            for word in words:
                candidate = f"{current} {word}".strip()
                if len(candidate) > max_chars and current:
                    rows.append(current)
                    current = word
                else:
                    current = candidate
            if current:
                rows.append(current)
            for row in rows:
                pdf.drawString(x, y, row)
                y -= 13
            y -= 6

    draw_column(left_lines, inch, height - inch - 58, 43)
    draw_column(right_lines, width / 2 + 12, height - inch - 58, 45)
    pdf.save()


def save_table_pdf(path: Path) -> None:
    profile = PROFILES["pl"]
    styles = pdf_styles()
    document = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    data = [
        [Paragraph(profile["summary_heading"], styles["body"]), Paragraph(profile["summary"], styles["body"])],
        [Paragraph(profile["experience_heading"], styles["body"]), Paragraph(" ".join(title + " " + " ".join(bullets) for title, bullets in profile["experience"]), styles["body"])],
        [Paragraph(profile["skills_heading"], styles["body"]), Paragraph(profile["skills"], styles["body"])],
        [Paragraph(profile["education_heading"], styles["body"]), Paragraph(profile["education"], styles["body"])],
    ]
    table = Table(data, colWidths=[1.875 * inch, 4.625 * inch], repeatRows=0)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAB4C0")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    document.build([Paragraph(profile["name"], styles["title"]), Paragraph(profile["title"], styles["subtitle"]), Spacer(1, 8), table])


def save_short_pdf(path: Path) -> None:
    styles = pdf_styles()
    document = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    document.build([Paragraph(line, styles["body"]) for line in SHORT_TEXT.split("\n")])


def save_scan_pdf(path: Path, scan_path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.drawImage(str(scan_path), 0, 0, width=letter[0], height=letter[1], preserveAspectRatio=True, anchor="c")
    pdf.save()


def save_blank_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.showPage()
    pdf.save()


def save_encrypted_pdf(source: Path, path: Path) -> None:
    reader = PdfReader(str(source))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.encrypt("synthetic-password")
    with path.open("wb") as stream:
        writer.write(stream)


def manifest_entry(number: int, filename: str, case: str, expected_code: str, expected_reason: str | None = None):
    return {
        "id": number,
        "file": filename,
        "format": Path(filename).suffix[1:],
        "case": case,
        "expected_code": expected_code,
        **({"expected_reason": expected_reason} if expected_reason else {}),
    }


def generate(outdir: Path) -> None:
    if outdir.exists():
        if any(outdir.iterdir()):
            raise SystemExit("Refusing to overwrite a non-empty output directory")
    else:
        outdir.mkdir(parents=True)
    scan_path = outdir / "scan-source.png"
    make_scan_image(scan_path)

    save_standard_docx(outdir / "01-docx-standard-pl.docx", "pl")
    save_standard_docx(outdir / "02-docx-standard-en.docx", "en")
    save_table_docx(outdir / "03-docx-table-layout.docx")
    save_two_column_docx(outdir / "04-docx-two-column.docx")
    save_header_contact_docx(outdir / "05-docx-header-contact.docx")
    save_mixed_language_docx(outdir / "06-docx-mixed-language.docx")
    save_short_docx(outdir / "07-docx-short-genuine.docx")
    save_image_only_docx(outdir / "08-docx-image-only.docx", scan_path)
    save_standard_docx(outdir / "_valid-for-corruption.docx", "pl")
    raw_docx = (outdir / "_valid-for-corruption.docx").read_bytes()
    (outdir / "09-docx-truncated.docx").write_bytes(raw_docx[: len(raw_docx) // 2])
    save_malformed_docx(outdir / "10-docx-malformed-xml.docx")

    save_standard_pdf(outdir / "11-pdf-standard-pl.pdf", "pl")
    save_standard_pdf(outdir / "12-pdf-standard-en.pdf", "en")
    save_multicolumn_pdf(outdir / "13-pdf-two-column.pdf")
    save_table_pdf(outdir / "14-pdf-table-layout.pdf")
    save_short_pdf(outdir / "15-pdf-short-genuine.pdf")
    save_scan_pdf(outdir / "16-pdf-scan-image-only.pdf", scan_path)
    save_blank_pdf(outdir / "17-pdf-blank.pdf")
    save_encrypted_pdf(outdir / "11-pdf-standard-pl.pdf", outdir / "18-pdf-encrypted.pdf")
    (outdir / "19-pdf-truncated.pdf").write_bytes(b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\nendobj\n")

    for scratch in (scan_path, outdir / "_valid-for-corruption.docx"):
        scratch.unlink()

    manifest = [
        manifest_entry(1, "01-docx-standard-pl.docx", "standard Polish text", "ACCEPTED"),
        manifest_entry(2, "02-docx-standard-en.docx", "standard English text", "ACCEPTED"),
        manifest_entry(3, "03-docx-table-layout.docx", "content stored in table cells", "ACCEPTED"),
        manifest_entry(4, "04-docx-two-column.docx", "two-column section", "ACCEPTED"),
        manifest_entry(5, "05-docx-header-contact.docx", "contact details in header", "ACCEPTED"),
        manifest_entry(6, "06-docx-mixed-language.docx", "mixed Polish and English", "ACCEPTED"),
        manifest_entry(7, "07-docx-short-genuine.docx", "very concise but genuine CV", "ACCEPTED"),
        manifest_entry(8, "08-docx-image-only.docx", "CV embedded only as an image", "UNREADABLE_DOCUMENT", "docx_no_text"),
        manifest_entry(9, "09-docx-truncated.docx", "truncated ZIP container", "INVALID_FILE_SIGNATURE"),
        manifest_entry(10, "10-docx-malformed-xml.docx", "valid ZIP with malformed document XML", "UNREADABLE_DOCUMENT", "docx_parse_failed"),
        manifest_entry(11, "11-pdf-standard-pl.pdf", "standard Polish text", "ACCEPTED"),
        manifest_entry(12, "12-pdf-standard-en.pdf", "standard English text", "ACCEPTED"),
        manifest_entry(13, "13-pdf-two-column.pdf", "two-column visual layout", "ACCEPTED"),
        manifest_entry(14, "14-pdf-table-layout.pdf", "content in a visual table", "ACCEPTED"),
        manifest_entry(15, "15-pdf-short-genuine.pdf", "very concise but genuine CV", "ACCEPTED"),
        manifest_entry(16, "16-pdf-scan-image-only.pdf", "image-only scanned CV", "UNREADABLE_DOCUMENT", "pdf_no_text_items"),
        manifest_entry(17, "17-pdf-blank.pdf", "blank PDF", "UNREADABLE_DOCUMENT", "pdf_no_text_items"),
        manifest_entry(18, "18-pdf-encrypted.pdf", "password-protected PDF", "UNREADABLE_DOCUMENT", "pdf_encrypted"),
        manifest_entry(19, "19-pdf-truncated.pdf", "truncated PDF structure", "UNREADABLE_DOCUMENT", "pdf_parse_failed"),
    ]
    (outdir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--outdir", required=True, type=Path)
    args = parser.parse_args()
    generate(args.outdir.resolve())
    print(json.dumps({"generated": 19, "outdir": str(args.outdir.resolve())}))


if __name__ == "__main__":
    font_path = Path(__file__).resolve().parents[2] / "src" / "assets" / "og-font.ttf"
    pdfmetrics.registerFont(TTFont("CVFixture", str(font_path)))
    main()
